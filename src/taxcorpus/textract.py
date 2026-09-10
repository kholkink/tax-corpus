"""Извлечение текста из файлов дела (docx, pdf, xlsx, md, txt, html).

docx — стандартной библиотекой (zipfile + XML), pdf — pypdf (только текстовый слой;
скан помечается как пустой), xlsx — openpyxl в markdown-таблицы. Результат
кэшируется рядом с индексом дела по sha256 файла (см. workspace.py).
"""

from __future__ import annotations

import html as _html
import io
import re
import zipfile
from pathlib import Path

TEXT_SUFFIXES = {".md", ".txt", ".csv", ".json"}


def docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    paragraphs = []
    for para in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        runs = re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", para, re.S)
        line = _html.unescape("".join(runs)).strip()
        if line:
            paragraphs.append(line)
    return "\n\n".join(paragraphs)


def pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("для pdf нужен pypdf: pip install -e '.[workspace]'") from exc
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(f"[стр. {i}]\n{text}" if text else f"[стр. {i}]\n(текстового слоя нет — нужен OCR)")
    return "\n\n".join(pages)


def xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("для xlsx нужен openpyxl: pip install -e '.[workspace]'") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            rows.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
            if len(rows) >= 500:
                rows.append("| … (лист обрезан до 500 строк) |")
                break
        parts.append(f"## Лист «{ws.title}»\n" + "\n".join(rows))
    return "\n\n".join(parts)


def html_text(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    text = re.sub(r"<(script|style).*?</\1>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<br\s*/?>|</p>|</div>|</tr>|</h\d>", "\n", text, flags=re.I)
    text = _html.unescape(re.sub(r"<[^>]+>", " ", text))
    return re.sub(r"\n\s*\n+", "\n\n", re.sub(r"[ \t\xa0]+", " ", text)).strip()


def extract_text(path: str | Path) -> str:
    """Файл -> текст. Неподдерживаемый формат -> RuntimeError."""
    p = Path(path)
    suffix = p.suffix.lower()
    data = p.read_bytes()
    if suffix in TEXT_SUFFIXES:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        return docx_text(data)
    if suffix == ".pdf":
        return pdf_text(data)
    if suffix in (".xlsx", ".xlsm"):
        return xlsx_text(data)
    if suffix in (".html", ".htm"):
        return html_text(data)
    raise RuntimeError(f"формат {suffix or 'без расширения'} не поддерживается "
                       "(docx, pdf, xlsx, md, txt, csv, json, html)")
