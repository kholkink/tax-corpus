"""Экспорт файла агента (markdown с шапкой провенанса) в DOCX (F7 плана ПО).

Поддержка: заголовки #..####, абзацы, списки (-, *, 1.), таблицы |a|b|, **жирный**, *курсив*,
цитаты >, разделители ---. Шапка провенанса (---…---) не печатается в теле: as_of, снимок
корпуса, версия и проверка цитат — в свойствах документа (comments) и на последней странице
«Провенанс». Секции агента <!-- agent: … --> в экспорт не попадают.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

RE_FRONT = re.compile(r"^---\n(.*?)\n---\n", re.S)
RE_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
RE_AGENT = re.compile(r"<!--\s*agent:.*?-->", re.S)
RE_COMMENT = re.compile(r"<!--.*?-->", re.S)


def parse_header(text: str) -> tuple[dict, str]:
    m = RE_FRONT.match(text)
    if not m:
        return {}, text
    head = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            try:
                head[k.strip()] = json.loads(v.strip())
            except json.JSONDecodeError:
                head[k.strip()] = v.strip()
    return head, text[m.end():]


def _add_inline(paragraph, text: str) -> None:
    for part in RE_INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def markdown_to_docx(text: str, reference_docx: str | Path | None = None, title: str | None = None) -> bytes:
    """markdown -> DOCX bytes. reference_docx — файл со стилями фирмы (Normal, Heading 1..3, List Bullet …)."""
    from docx import Document
    from docx.shared import Pt

    head, body = parse_header(text)
    body = RE_AGENT.sub("", body)
    body = RE_COMMENT.sub("", body)
    doc = Document(str(reference_docx)) if reference_docx else Document()
    if reference_docx:
        for p in list(doc.paragraphs):  # эталонный файл нужен только ради стилей
            p._element.getparent().remove(p._element)
    else:
        doc.styles["Normal"].font.name = "Times New Roman"
        doc.styles["Normal"].font.size = Pt(12)

    lines = body.splitlines()
    i = 0
    first_heading = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.match(r"^[\s:|-]+$", lines[i].strip().strip("|")):
                    rows.append(cells)
                i += 1
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_inline(cell.paragraphs[0], cells[c] if c < len(cells) else "")
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            h = doc.add_heading("", level=min(level, 4))
            _add_inline(h, m.group(2).strip())
            if first_heading is None:
                first_heading = m.group(2).strip()
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\s*\d+[.)]\s+", "", line))
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote" if "Intense Quote" in [s.name for s in doc.styles] else None)
            _add_inline(p, line.lstrip("> ").strip())
            i += 1
            continue
        if re.match(r"^-{3,}$|^\*{3,}$", line.strip()):
            doc.add_paragraph("")
            i += 1
            continue
        # абзац: соседние непустые строки склеиваются
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\s*[-*]\s|\s*\d+[.)]\s|\||>|-{3,})", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        _add_inline(p, " ".join(buf))

    doc.core_properties.title = title or first_heading or "Документ tax-corpus"
    if head:
        brief = {k: head[k] for k in ("as_of", "corpus_snapshot", "version", "template") if k in head}
        doc.core_properties.comments = json.dumps(brief, ensure_ascii=False)[:250]   # лимит свойства — 255 символов
        doc.core_properties.subject = f"as_of {head.get('as_of', '')} · снимок корпуса {head.get('corpus_snapshot', '')}"
        doc.add_page_break()
        doc.add_heading("Провенанс", level=2)
        for key in ("as_of", "corpus_snapshot", "version", "written_at", "template", "summary"):
            if head.get(key) not in (None, ""):
                doc.add_paragraph(f"{key}: {head[key]}")
        sources = head.get("sources") or []
        if sources:
            doc.add_paragraph("Источники (единицы корпуса): " + ", ".join(sources))
        ver = head.get("verification")
        if isinstance(ver, dict):
            doc.add_paragraph("Проверка цитат: " + ("замечаний нет" if ver.get("ok") else "; ".join(ver.get("problems") or [])))
    buf_out = io.BytesIO()
    doc.save(buf_out)
    return buf_out.getvalue()


def export_workspace_file(ws, rel: str, reference_docx: str | Path | None = None) -> bytes:
    p = ws._resolve(rel)
    if not p.is_file():
        raise FileNotFoundError(rel)
    return markdown_to_docx(p.read_text(encoding="utf-8"), reference_docx)
